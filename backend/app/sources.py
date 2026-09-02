from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import Optional

import pdfplumber

ACCEPTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".csv", ".md"}


def extract_text(path: Path) -> tuple[str, Optional[str]]:
    """
    Extract plain text from a file.
    Returns (text, warning_or_none).
    """
    ext = path.suffix.lower()

    if ext == ".pdf":
        return _extract_pdf(path)

    if ext == ".docx":
        return _extract_docx(path)

    if ext in (".txt", ".md"):
        text = path.read_text(encoding="utf-8", errors="replace").strip()
        return text, None

    if ext == ".csv":
        return _extract_csv(path)

    return "", f"Unsupported file type '{ext}' — skipped."


def _extract_pdf(path: Path) -> tuple[str, Optional[str]]:
    pages: list[str] = []
    try:
        with pdfplumber.open(str(path)) as pdf:
            for i, page in enumerate(pdf.pages, start=1):
                text = page.extract_text() or ""
                pages.append(f"--- page {i} ---\n\n{text.strip()}")
    except Exception as exc:
        return "", f"Could not read PDF: {exc}"

    full = "\n\n".join(pages).strip()
    if len(full) < 50:
        return "", (
            "no text layer — probably a scan. OCR is not supported in v0."
        )
    return full, None


def _extract_docx(path: Path) -> tuple[str, Optional[str]]:
    try:
        import docx  # python-docx
    except ImportError:
        return "", "python-docx not installed."

    try:
        doc = docx.Document(str(path))
    except Exception as exc:
        return "", f"Could not read DOCX: {exc}"

    lines: list[str] = []

    # Paragraphs
    for para in doc.paragraphs:
        t = para.text.strip()
        if t:
            lines.append(t)

    # Tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                t = cell.text.strip()
                if t:
                    lines.append(t)

    return "\n".join(lines), None


def _extract_csv(path: Path) -> tuple[str, Optional[str]]:
    try:
        raw = path.read_text(encoding="utf-8", errors="replace")
        reader = csv.DictReader(io.StringIO(raw))
        lines: list[str] = []
        for row in reader:
            parts = [f"{k}: {v}" for k, v in row.items() if v.strip()]
            if parts:
                lines.append(", ".join(parts))
        return "\n".join(lines), None
    except Exception as exc:
        return "", f"Could not read CSV: {exc}"
